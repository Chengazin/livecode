import json
import subprocess
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
OUT_PATH = ROOT / "economic_report_livecode.docx"


@dataclass(frozen=True)
class BackendCounts:
    if_count: int
    elseif_count: int
    for_count: int
    foreach_count: int
    while_count: int
    api_routes_count: int
    gateway_routes_count: int


@dataclass(frozen=True)
class FrontendCounts:
    pages_count: int
    components_count: int


def iter_code_files(paths: Iterable[Path], extensions: set[str]) -> Iterable[Path]:
    for base in paths:
        if not base.exists():
            continue
        for path in base.rglob("*"):
            if path.is_file() and path.suffix.lower() in extensions:
                yield path


def count_keyword_occurrences(files: Iterable[Path], keyword: str) -> int:
    # Simple, reproducible count. Intentionally does not try to parse language grammar.
    # We rely on word-boundaries to avoid matching identifiers like "diff" or "before".
    import re

    pattern = re.compile(rf"\b{re.escape(keyword)}\b")
    total = 0
    for path in files:
        try:
            text = path.read_text(encoding="utf-8", errors="ignore")
        except OSError:
            continue
        total += len(pattern.findall(text))
    return total


def load_laravel_routes_json() -> list[dict]:
    proc = subprocess.run(
        ["php", "artisan", "route:list", "--json"],
        cwd=ROOT,
        capture_output=True,
        check=True,
    )

    # On Windows, artisan often prints UTF-16LE with BOM.
    raw = proc.stdout
    try:
        decoded = raw.decode("utf-16")
    except UnicodeDecodeError:
        decoded = raw.decode("utf-8", errors="replace")

    return json.loads(decoded)


def count_api_routes() -> int:
    routes = load_laravel_routes_json()
    return sum(1 for route in routes if str(route.get("uri", "")).startswith("api/"))


def compute_backend_counts() -> BackendCounts:
    backend_php_dirs = [ROOT / "app", ROOT / "routes"]
    gateway_js_dirs = [ROOT / "terminal-gateway" / "src"]

    php_files = list(iter_code_files(backend_php_dirs, {".php"}))
    js_files = list(iter_code_files(gateway_js_dirs, {".js"}))
    backend_files = php_files + js_files

    if_count = count_keyword_occurrences(backend_files, "if")
    elseif_count = count_keyword_occurrences(backend_files, "elseif")
    for_count = count_keyword_occurrences(backend_files, "for")
    foreach_count = count_keyword_occurrences(backend_files, "foreach")
    while_count = count_keyword_occurrences(backend_files, "while")
    api_routes_count = count_api_routes()

    # terminal-gateway defines:
    # - HTTP: GET /healthz
    # - WebSocket upgrade path (default: /terminal)
    gateway_routes_count = 2

    return BackendCounts(
        if_count=if_count,
        elseif_count=elseif_count,
        for_count=for_count,
        foreach_count=foreach_count,
        while_count=while_count,
        api_routes_count=api_routes_count,
        gateway_routes_count=gateway_routes_count,
    )


def compute_frontend_counts() -> FrontendCounts:
    pages_dir = ROOT / "frontend" / "src" / "pages"
    components_dir = ROOT / "frontend" / "src" / "components"

    pages = list(iter_code_files([pages_dir], {".vue"}))
    components = list(iter_code_files([components_dir], {".vue"}))

    return FrontendCounts(pages_count=len(pages), components_count=len(components))


def round_hours(value: float) -> int:
    return int(round(value))


def build_report(backend: BackendCounts, frontend: FrontendCounts) -> Document:
    # Method (explicitly stated in the user's requirements)
    conditionals = backend.if_count + backend.elseif_count
    loops = backend.for_count + backend.foreach_count + backend.while_count
    routes = backend.api_routes_count + backend.gateway_routes_count
    ui_sections = frontend.pages_count + frontend.components_count
    q_blocks = conditionals + loops + routes + ui_sections

    # Coefficients (can be tuned; defaults reflect the provided sample doc)
    c_complexity = 1.08
    p_correction = 0.08

    q_adjusted = q_blocks * c_complexity * (1 + p_correction)

    # Labor distribution coefficients inferred from the sample:
    # total ~ 0.699 * Q
    k_top = 0.01592  # task preparation & description
    k_research = 0.02817  # research of algorithms
    k_algorithms = 0.07584  # flowcharts / algorithms
    k_coding = 0.07584  # coding
    k_debug = 0.37038  # debugging
    k_docs = 0.13265  # documentation

    t_top = round_hours(q_adjusted * k_top)
    t_research = round_hours(q_adjusted * k_research)
    t_algorithms = round_hours(q_adjusted * k_algorithms)
    t_coding = round_hours(q_adjusted * k_coding)
    t_debug = round_hours(q_adjusted * k_debug)
    t_docs = round_hours(q_adjusted * k_docs)
    t_total = t_top + t_research + t_algorithms + t_coding + t_debug + t_docs

    # Work calendar assumptions (same structure as the sample)
    hours_per_day = 6
    team_size = 4
    days_per_month = 22

    days_single = t_total / hours_per_day
    days_team = days_single / team_size
    months_team = days_team / days_per_month

    # Cost assumptions (student project defaults; adjust if needed)
    hourly_rate_rub = 250.0
    extra_salary_share = 0.10
    social_tax_share = 0.30
    overhead_share = 0.60

    pc_power_kw = 0.25
    electricity_tariff_rub_per_kwh = 5.85

    pc_cost_rub = 60000.0
    depreciation_years = 5
    study_months_per_year = 10

    base_salary = t_total * hourly_rate_rub
    extra_salary = base_salary * extra_salary_share
    total_salary = base_salary + extra_salary
    social_tax = total_salary * social_tax_share
    overhead = total_salary * overhead_share

    electricity_kwh = pc_power_kw * t_total
    electricity_cost = electricity_kwh * electricity_tariff_rub_per_kwh

    depreciation_hours_total = depreciation_years * study_months_per_year * days_per_month * hours_per_day
    depreciation_per_hour = pc_cost_rub / depreciation_hours_total if depreciation_hours_total else 0.0
    depreciation_cost = depreciation_per_hour * t_total

    cost_total = total_salary + social_tax + overhead + electricity_cost + depreciation_cost

    profit_share = 0.20
    vat_share = 0.20
    price_no_vat = cost_total * (1 + profit_share)
    price_with_vat = price_no_vat * (1 + vat_share)

    doc = Document()

    doc.add_heading("РАСЧЕТ ЭКОНОМИЧЕСКОЙ ЧАСТИ ПРОГРАММНОГО ПРОДУКТА", level=0)
    doc.add_paragraph("Веб-приложение «LiveCode» (Laravel 12 API + Vue 3 SPA + Node.js terminal-gateway).")

    doc.add_heading("1. Краткое описание проекта", level=1)
    doc.add_paragraph(
        "LiveCode — веб-приложение для работы с проектами и исходным кодом в браузере: "
        "аутентификация, управление проектами, файловая система, редактор кода, "
        "реалтайм-синхронизация/чат, терминальные сессии, интеграция с Forgejo, админ-панель."
    )

    doc.add_heading("2. Методика расчета объема работ", level=1)
    doc.add_paragraph(
        "В отличие от классического подхода (подсчет строк кода), в данном расчете используется "
        "подсчет базовых блоков."
    )
    doc.add_paragraph(
        "Backend: условные блоки (if/elseif), циклы (for/foreach/while), количество маршрутов API."
    )
    doc.add_paragraph(
        "Frontend: уникальные пользовательские секции (страницы и крупные UI-компоненты). "
        "Повторное использование визуальных блоков в цену не включается."
    )

    doc.add_heading("3. Результаты подсчета базовых блоков", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Категория"
    hdr[1].text = "Ед."
    hdr[2].text = "Количество"

    def add_row(category: str, unit: str, value: int) -> None:
        row = table.add_row().cells
        row[0].text = category
        row[1].text = unit
        row[2].text = str(value)

    add_row("Backend: if", "блок", backend.if_count)
    add_row("Backend: elseif", "блок", backend.elseif_count)
    add_row("Backend: циклы (for/foreach/while)", "блок", loops)
    add_row("Backend: маршруты Laravel API (api/*)", "роут", backend.api_routes_count)
    add_row("Backend: маршруты terminal-gateway (/healthz + WS path)", "роут", backend.gateway_routes_count)
    add_row("Frontend: уникальные страницы (frontend/src/pages)", "секция", frontend.pages_count)
    add_row("Frontend: уникальные UI-секции (frontend/src/components)", "секция", frontend.components_count)
    add_row("Итого q (базовые блоки)", "ед.", q_blocks)

    doc.add_heading("4. Расчет трудоемкости", level=1)
    doc.add_paragraph(f"Коэффициент сложности c = {c_complexity:.2f}.")
    doc.add_paragraph(f"Коэффициент коррекции p = {p_correction:.2f}.")
    doc.add_paragraph(f"Q = q * c * (1 + p) = {q_blocks} * {c_complexity:.2f} * {1 + p_correction:.2f} = {q_adjusted:.2f}.")

    labor = doc.add_table(rows=1, cols=3)
    labor.style = "Table Grid"
    hdr2 = labor.rows[0].cells
    hdr2[0].text = "Этап"
    hdr2[1].text = "Формула"
    hdr2[2].text = "Трудоемкость, чел.-час"

    def add_labor_row(name: str, formula: str, hours: int) -> None:
        row = labor.add_row().cells
        row[0].text = name
        row[1].text = formula
        row[2].text = str(hours)

    add_labor_row("Подготовка/описание задачи", f"{k_top:.5f} * Q", t_top)
    add_labor_row("Исследование алгоритмов", f"{k_research:.5f} * Q", t_research)
    add_labor_row("Разработка алгоритмов/схем", f"{k_algorithms:.5f} * Q", t_algorithms)
    add_labor_row("Программирование", f"{k_coding:.5f} * Q", t_coding)
    add_labor_row("Отладка", f"{k_debug:.5f} * Q", t_debug)
    add_labor_row("Документирование", f"{k_docs:.5f} * Q", t_docs)
    add_labor_row("Итого", "Σ", t_total)

    doc.add_paragraph(
        f"Перевод в дни при {hours_per_day} час/день: {t_total} / {hours_per_day} = {days_single:.1f} дней (1 человек)."
    )
    doc.add_paragraph(
        f"Для команды из {team_size} человек: {days_single:.1f} / {team_size} = {days_team:.1f} дней."
    )
    doc.add_paragraph(f"В месяцах (при {days_per_month} днях/мес): {days_team:.1f} / {days_per_month} = {months_team:.2f} мес.")

    doc.add_heading("5. Расчет себестоимости и цены", level=1)
    doc.add_paragraph(
        "Далее приведен пример расчета себестоимости для студенческого проекта. "
        "При необходимости измените ставку, проценты накладных расходов и коэффициенты."
    )

    cost = doc.add_table(rows=1, cols=3)
    cost.style = "Table Grid"
    hdr3 = cost.rows[0].cells
    hdr3[0].text = "Статья затрат"
    hdr3[1].text = "Расчет"
    hdr3[2].text = "Сумма, руб."

    def add_cost_row(name: str, formula: str, value: float) -> None:
        row = cost.add_row().cells
        row[0].text = name
        row[1].text = formula
        row[2].text = f"{value:,.0f}".replace(",", " ")

    add_cost_row("Основная зарплата", f"{t_total} * {hourly_rate_rub:.0f}", base_salary)
    add_cost_row("Доп. зарплата", f"{extra_salary_share:.0%} от осн.", extra_salary)
    add_cost_row("Итого зарплата", "осн. + доп.", total_salary)
    add_cost_row("Страховые взносы", f"{social_tax_share:.0%} от зарплаты", social_tax)
    add_cost_row("Накладные расходы", f"{overhead_share:.0%} от зарплаты", overhead)
    add_cost_row(
        "Электроэнергия",
        f"{pc_power_kw:.2f} кВт * {t_total} ч * {electricity_tariff_rub_per_kwh:.2f}",
        electricity_cost,
    )
    add_cost_row(
        "Амортизация ПК",
        f"{pc_cost_rub:.0f} / ({depreciation_years} лет * {study_months_per_year} мес * {days_per_month} дн * {hours_per_day} ч) * {t_total} ч",
        depreciation_cost,
    )

    add_cost_row("Себестоимость", "Σ", cost_total)
    add_cost_row("Прибыль", f"{profit_share:.0%} от себестоимости", cost_total * profit_share)
    add_cost_row("Цена без НДС", "себестоимость + прибыль", price_no_vat)
    add_cost_row("НДС", f"{vat_share:.0%} от цены без НДС", price_no_vat * vat_share)
    add_cost_row("Итоговая цена", "цена без НДС + НДС", price_with_vat)

    doc.add_paragraph("Примечание: расчеты выполнены автоматически на основе текущего состояния репозитория.")

    return doc


def main() -> None:
    backend = compute_backend_counts()
    frontend = compute_frontend_counts()
    doc = build_report(backend, frontend)
    doc.save(OUT_PATH)
    print(f"Saved: {OUT_PATH}")


if __name__ == "__main__":
    main()
